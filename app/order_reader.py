from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class OrderReadError(RuntimeError):
    pass


COUNTRY_SUFFIX_ALIASES = {
    "bulgaria": "BUŁGARIA",
    "bulgaria bg": "BUŁGARIA",
    "czechy": "CZECHY",
    "republika czeska": "CZECHY",
    "czech republic": "CZECHY",
    "finlandia": "FINLANDIA",
    "litwa": "LITWA",
    "portugalia": "PORTUGALIA",
    "rumunia": "RUMUNIA",
    "romania": "RUMUNIA",
    "slowenia": "SŁOWENIA",
    "slovenia": "SŁOWENIA",
    "wegry": "WĘGRY",
    "hungary": "WĘGRY",
    "wlochy": "WŁOCHY",
    "italia": "WŁOCHY",
    "italy": "WŁOCHY",
}


PRODUCT_HEADER_ALIASES = {
    "towarnazwa",
    "nazwatowaru",
    "nazwaproduktu",
    "produkt",
    "productname",
    "nazwa",
}


QUANTITY_HEADER_ALIASES = {
    "ilosc",
    "liczba",
    "liczbasztuk",
    "quantity",
    "qty",
}


CATALOG_HEADER_ALIASES = {
    "towarnumerkatalogowy",
    "numerkatalogowy",
    "nrkatalogowy",
    "catalognumber",
    "sku",
}


CAPACITY_PATTERN = re.compile(
    r"(?<!\d)"
    r"(\d+(?:[.,]\d+)?)"
    r"\s*(ml|l)"
    r"\b",
    flags=re.IGNORECASE,
)


CAPACITY_AT_END_PATTERN = re.compile(
    r"\s+"
    r"\d+(?:[.,]\d+)?"
    r"\s*(?:ml|l)"
    r"\s*$",
    flags=re.IGNORECASE,
)


KIT_OR_SET_PATTERN = re.compile(
    r"\b(?:KIT|SET)\b",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class OrderItem:
    row_number: int
    product_name: str
    quantity: int
    capacity_liters: float | None
    label_format: str | None
    product_folder_name: str
    catalog_number: str | None = None


@dataclass(frozen=True)
class ParsedOrder:
    source_path: Path
    country: str
    worksheet_name: str
    items: list[OrderItem]


def _normalize_text(value: Any) -> str:
    text = (
        str(value or "")
        .strip()
        .casefold()
    )

    decomposed = unicodedata.normalize(
        "NFKD",
        text,
    )

    without_accents = "".join(
        character
        for character in decomposed
        if not unicodedata.combining(
            character
        )
    )

    return re.sub(
        r"[^a-z0-9]+",
        " ",
        without_accents,
    ).strip()


def _normalize_header(value: Any) -> str:
    return _normalize_text(
        value
    ).replace(" ", "")


def extract_country(
    file_path: str | Path,
) -> str:
    filename = Path(file_path).stem

    normalized_filename = (
        _normalize_text(filename)
    )

    aliases = sorted(
        COUNTRY_SUFFIX_ALIASES.items(),
        key=lambda item: len(item[0]),
        reverse=True,
    )

    for alias, canonical_country in aliases:
        if (
            normalized_filename == alias
            or normalized_filename.endswith(
                f" {alias}"
            )
        ):
            return canonical_country

    supported_countries = (
        "BUŁGARIA, CZECHY, FINLANDIA, "
        "LITWA, PORTUGALIA, RUMUNIA, "
        "SŁOWENIA, WĘGRY lub WŁOCHY"
    )

    raise OrderReadError(
        "Nie udało się rozpoznać państwa "
        "z końca nazwy pliku. "
        "Nazwa pliku musi kończyć się nazwą "
        "państwa, np. „Ahifi CZECHY.xlsx”. "
        f"Obsługiwane państwa: "
        f"{supported_countries}."
    )


def extract_capacity_liters(
    product_name: str,
) -> float | None:
    matches = list(
        CAPACITY_PATTERN.finditer(
            product_name
        )
    )

    if not matches:
        return None

    match = matches[-1]

    numeric_value = (
        match.group(1)
        .replace(",", ".")
    )

    try:
        capacity = float(
            numeric_value
        )
    except ValueError:
        return None

    unit = match.group(2).casefold()

    if unit == "ml":
        capacity = capacity / 1000

    return capacity


def select_label_format(
    capacity_liters: float | None,
    product_name: str = "",
) -> str | None:
    # Produkty zawierające osobne słowo
    # KIT albo SET zawsze otrzymują format
    # 45x110, nawet jeśli nazwa nie zawiera
    # informacji o pojemności.
    if KIT_OR_SET_PATTERN.search(
        product_name
    ):
        return "45x110"

    if capacity_liters is None:
        return None

    small_capacities = (
        0.2,
        0.5,
    )

    large_capacities = (
        1.0,
        5.0,
        10.0,
    )

    if any(
        math.isclose(
            capacity_liters,
            expected,
            abs_tol=0.001,
        )
        for expected in small_capacities
    ):
        return "45x45"

    if any(
        math.isclose(
            capacity_liters,
            expected,
            abs_tol=0.001,
        )
        for expected in large_capacities
    ):
        return "45x110"

    return None


def extract_product_folder_name(
    product_name: str,
) -> str:
    name = str(
        product_name
    ).strip()

    name = re.sub(
        r"^\s*ADBL\b[\s:-]*",
        "",
        name,
        flags=re.IGNORECASE,
    )

    name = CAPACITY_AT_END_PATTERN.sub(
        "",
        name,
    )

    name = re.sub(
        r"\s+",
        " ",
        name,
    )

    return name.strip(
        " -–—"
    )


def parse_quantity(value: Any) -> int:
    if value is None:
        raise OrderReadError(
            "Brak ilości produktu."
        )

    if isinstance(value, bool):
        raise OrderReadError(
            "Nieprawidłowa ilość produktu."
        )

    if isinstance(value, int):
        quantity = value

    elif isinstance(value, float):
        if not value.is_integer():
            raise OrderReadError(
                "Ilość musi być liczbą "
                f"całkowitą: {value}."
            )

        quantity = int(value)

    else:
        text = str(value).strip()

        if not text:
            raise OrderReadError(
                "Brak ilości produktu."
            )

        text = text.replace(
            " ",
            "",
        )

        text = text.replace(
            ",",
            ".",
        )

        try:
            numeric_value = float(text)
        except ValueError as exc:
            raise OrderReadError(
                f"Nieprawidłowa ilość: "
                f"{value}."
            ) from exc

        if not numeric_value.is_integer():
            raise OrderReadError(
                "Ilość musi być liczbą "
                f"całkowitą: {value}."
            )

        quantity = int(
            numeric_value
        )

    if quantity < 0:
        raise OrderReadError(
            "Ilość nie może być ujemna: "
            f"{quantity}."
        )

    return quantity


def _find_column(
    header_values: tuple[Any, ...],
    aliases: set[str],
) -> int | None:
    for column_index, value in enumerate(
        header_values
    ):
        normalized = (
            _normalize_header(value)
        )

        if normalized in aliases:
            return column_index

    return None


def _optional_text(
    value: Any,
) -> str | None:
    if value is None:
        return None

    text = str(value).strip()

    return text or None


def read_order(
    file_path: str | Path,
) -> ParsedOrder:
    source_path = Path(file_path)

    if not source_path.is_file():
        raise OrderReadError(
            "Nie znaleziono pliku "
            f"zamówienia: {source_path}"
        )

    country = extract_country(
        source_path
    )

    try:
        workbook = load_workbook(
            source_path,
            read_only=True,
            data_only=True,
        )

    except Exception as exc:
        raise OrderReadError(
            "Nie udało się otworzyć "
            f"pliku Excel: {exc}"
        ) from exc

    try:
        worksheet = workbook.active

        rows = worksheet.iter_rows(
            values_only=True
        )

        header_row_number: (
            int | None
        ) = None

        product_column: (
            int | None
        ) = None

        quantity_column: (
            int | None
        ) = None

        catalog_column: (
            int | None
        ) = None

        for row_number, row in enumerate(
            rows,
            start=1,
        ):
            row_values = tuple(row)

            possible_product_column = (
                _find_column(
                    row_values,
                    PRODUCT_HEADER_ALIASES,
                )
            )

            possible_quantity_column = (
                _find_column(
                    row_values,
                    QUANTITY_HEADER_ALIASES,
                )
            )

            if (
                possible_product_column is None
                or possible_quantity_column
                is None
            ):
                continue

            header_row_number = row_number

            product_column = (
                possible_product_column
            )

            quantity_column = (
                possible_quantity_column
            )

            catalog_column = _find_column(
                row_values,
                CATALOG_HEADER_ALIASES,
            )

            break

        if header_row_number is None:
            raise OrderReadError(
                "W arkuszu nie znaleziono "
                "wymaganych kolumn. "
                "Wymagane są tylko kolumny: "
                "nazwa produktu oraz ilość."
            )

        if product_column is None:
            raise OrderReadError(
                "Nie znaleziono kolumny "
                "z nazwą produktu."
            )

        if quantity_column is None:
            raise OrderReadError(
                "Nie znaleziono kolumny "
                "z ilością."
            )

        items: list[OrderItem] = []

        for row_number, row in enumerate(
            rows,
            start=header_row_number + 1,
        ):
            row_values = tuple(row)

            if (
                product_column
                >= len(row_values)
            ):
                continue

            product_value = row_values[
                product_column
            ]

            product_name = _optional_text(
                product_value
            )

            if not product_name:
                continue

            if (
                quantity_column
                >= len(row_values)
            ):
                raise OrderReadError(
                    "Brak ilości w wierszu "
                    f"{row_number}."
                )

            try:
                quantity = parse_quantity(
                    row_values[
                        quantity_column
                    ]
                )

            except OrderReadError as exc:
                raise OrderReadError(
                    f"Wiersz {row_number}, "
                    f"produkt „{product_name}”: "
                    f"{exc}"
                ) from exc

            # Pozycje z ilością 0 nie
            # wymagają żadnych etykiet.
            if quantity == 0:
                continue

            catalog_number = None

            if (
                catalog_column is not None
                and catalog_column
                < len(row_values)
            ):
                catalog_number = (
                    _optional_text(
                        row_values[
                            catalog_column
                        ]
                    )
                )

            capacity_liters = (
                extract_capacity_liters(
                    product_name
                )
            )

            label_format = (
                select_label_format(
                    capacity_liters,
                    product_name,
                )
            )

            product_folder_name = (
                extract_product_folder_name(
                    product_name
                )
            )

            items.append(
                OrderItem(
                    row_number=row_number,
                    product_name=product_name,
                    quantity=quantity,
                    capacity_liters=(
                        capacity_liters
                    ),
                    label_format=(
                        label_format
                    ),
                    product_folder_name=(
                        product_folder_name
                    ),
                    catalog_number=(
                        catalog_number
                    ),
                )
            )

        if not items:
            raise OrderReadError(
                "W zamówieniu nie znaleziono "
                "żadnych pozycji z ilością "
                "większą od zera."
            )

        return ParsedOrder(
            source_path=source_path,
            country=country,
            worksheet_name=worksheet.title,
            items=items,
        )

    finally:
        workbook.close()